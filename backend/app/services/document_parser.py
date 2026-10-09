"""
多格式文档解析器。

这是当前上传入库主链的统一解析入口。
这里通过注册表把 PDF、DOCX、PPTX、Markdown、HTML、TXT、LaTeX
等格式统一映射到相同的 `List[Chunk]` 契约。

关键边界：
  - 这里只负责"把文件解析成 Chunk"
  - 不负责写业务库、向量库或 LightRAG
  - 本地文档和 MinerU 产物共享同一套来源与分块契约
"""

from __future__ import annotations

# 抽象基类用于统一约束各格式解析器。
from abc import ABC, abstractmethod
from dataclasses import dataclass
# lru_cache 用于缓存 tokenizer，避免每次解析重复初始化。
from functools import lru_cache
import hashlib
from importlib.metadata import version
# Path 用来处理文件扩展名和读文件。
from pathlib import Path
# 正则用于 Markdown / LaTeX / 句子切分等文本拆分逻辑。
import re
# `Iterable` 给表格行抽象，`Optional` 用于单例注册表缓冲。
from typing import Any, Callable, Iterable, List, Literal, Optional

# 读取统一的分块配置。
from app.config import get_settings
# `Chunk` 是上传解析后统一返回的目标结构。
from app.models.models import Chunk
from app.models.knowledge_access_contracts import CONTENT_SCHEMA_VERSION, SourceFragment
from app.services.hf_local import configure_offline_env, resolve_local_model_source
from app.services.chunk_tokenizer import ApiTokenCounter, load_api_token_counter

# 全局配置里包含 parent / child chunk 大小等参数。
settings = get_settings()
# 默认解析注册表单例。
_REGISTRY: Optional["ParserRegistry"] = None
# 纯文本文件的解码回退顺序。
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "latin-1")
_TOKEN_SPLITTER_SEPARATOR = " "
_PARAGRAPH_SEPARATOR = "\n\n"
_MARKDOWN_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_MARKDOWN_PAGE_MARKER = re.compile(r"^<!--\s*page\s*:\s*(\d+)\s*-->$", re.IGNORECASE)
_LIST_ITEM_PATTERN = re.compile(r"^(?:[-*+]\s+|\d+\.\s+)")
_TABLE_ROW_PATTERN = re.compile(r"^\|.+\|\s*$")
# 公式片段整体受保护：`_` 在 LaTeX 里是下标，不能按 Markdown 强调符删掉。
# 捕获组是给 split 用的，奇数下标即公式片段。
_FORMULA_SPAN_PATTERN = re.compile(r"(?<![\\$])(\$\$(?:\\.|(?!\$\$)[^\\])*?\$\$|\$(?!\$)(?:\\.|[^\\$])+\$)", re.DOTALL)
_MARKDOWN_MARKERS_PATTERN = re.compile(r"[*_~`#>|]")
# 表格的索引文本与阅读正文都保留行列边界，避免数值失去对应列
_TABLE_MARKDOWN_MARKERS_PATTERN = re.compile(r"[*_~`#>]")
_TABLE_CELL_TAGS = ("td", "th")
_FORMAL_NUMBERED_HEADING_PATTERN = re.compile(r"^(?:\d+(?:\.\d+)*\b|\(\d+\)\s+)")
_ENUMERATED_UNIT_PATTERN = re.compile(r"(?<!\w)(\([a-zA-Z0-9]+\)|[A-Za-z]\)|\d+\))\s+(?=\S)")
_STRUCTURAL_UNIT_PATTERN = re.compile(r"^(?:\([a-zA-Z0-9]+\)|[A-Za-z]\)|\d+\)|\d+\.|[•\-*+])\s+")
_FORMULA_LABEL_PATTERN = re.compile(r"(?:\b[A-Z]{2,}\b|error|metric).{0,12}:\s*$", re.IGNORECASE)
_TABLE_CAPTION_PATTERN = re.compile(r"^table\s+\d+\b", re.IGNORECASE)
_EXCLUDED_SECTION_PATTERN = re.compile(
    r"\b(?:references|bibliography|acknowledg(?:e)?ments?|copyright|author contributions?)\b",
    re.IGNORECASE,
)
_MINERU_NOISE_TYPES = {
    "discarded",
    "header",
    "footer",
    "page_no",
    "page_number",
    "page_footnote",
    "footnote",
    "aside_text",
    "seal",
}
_ATOMIC_BLOCK_TYPES = {"table", "formula", "figure_caption", "code"}
_ERROR_MISSING_SECTION_PATH = "missing_section_path"
_ERROR_SINGLE_SECTION_KEY = "single_section_key"
_ERROR_STRUCTURED_BLOCK_TOO_LARGE = "structured_block_too_large"
_ERROR_SENTENCE_TOO_LARGE = "sentence_too_large"
_WARNING_ATOMIC_BLOCK_EXCEEDS_PARENT_LIMIT = "atomic_block_exceeds_parent_limit"
_GENERIC_SECTION_PATHS = {"untitled", "unknown", "content", "body", "text"}
_CHILD_MAX_MULTIPLIER = 1.5
_CHILD_SHORT_UNIT_RATIO = 0.25
_CHILD_SHORT_UNIT_TOKEN_LIMIT = 4


@dataclass(frozen=True)
class _MappedText:
    """规范阅读正文及同步传播的来源区间；切片参数和返回区间均为 Unicode 字符位置。"""

    text: str
    fragments: tuple[SourceFragment, ...] = ()

    def slice(self, start: int, end: int) -> "_MappedText":
        """参数为当前正文半开区间；返回正文切片，并同步裁剪来源项目和输出区间。"""
        fragments: list[SourceFragment] = []
        for fragment in self.fragments:
            left = max(start, fragment["text_start"])
            right = min(end, fragment["text_end"])
            if left >= right:
                continue
            source_start = fragment["source_start"] + left - fragment["text_start"]
            fragments.append({**fragment, "source_start": source_start,
                "source_end": source_start + right - left,
                "text_start": left - start, "text_end": right - start})
        return _MappedText(self.text[start:end], tuple(fragments))

    def strip(self) -> "_MappedText":
        """无参数；返回去除首尾空白后的正文，来源偏移随正文同步移动。"""
        start = len(self.text) - len(self.text.lstrip())
        return self.slice(start, max(start, len(self.text.rstrip())))


def _join_mapped(parts: Iterable[_MappedText], separator: str) -> _MappedText:
    """参数为有序正文片段和展示分隔符；返回合并正文，新增分隔符不冒充来源字符。"""
    texts: list[str] = []
    fragments: list[SourceFragment] = []
    offset = 0
    for part in parts:
        if not part.text:
            continue
        if texts:
            offset += len(separator)
        texts.append(part.text)
        fragments.extend({**fragment, "text_start": fragment["text_start"] + offset,
            "text_end": fragment["text_end"] + offset} for fragment in part.fragments)
        offset += len(part.text)
    return _MappedText(separator.join(texts), tuple(fragments))


def _split_mapped(value: _MappedText, pattern: str) -> list[_MappedText]:
    """参数为带来源正文和分隔正则；返回按实际匹配坐标切开的非空正文，不反查重复文字。"""
    parts = []
    cursor = 0
    for match in _unprotected_matches(value.text, pattern):
        part = value.slice(cursor, match.start()).strip()
        if part.text:
            parts.append(part)
        cursor = match.end()
    tail = value.slice(cursor, len(value.text)).strip()
    if tail.text:
        parts.append(tail)
    return parts


def _unprotected_matches(text: str, pattern: str | re.Pattern[str]) -> Iterable[re.Match[str]]:
    """参数为正文和自然边界模式；只返回公式之外的匹配坐标，保护公式内空行、句点与枚举。"""
    formulas = iter(_FORMULA_SPAN_PATTERN.finditer(text))
    formula = next(formulas, None)
    for match in re.finditer(pattern, text):
        while formula is not None and formula.end() <= match.start():
            formula = next(formulas, None)
        if formula is not None and match.start() < formula.end() and match.end() > formula.start():
            continue
        yield match

# 这一层的重要价值，是把"不同文件格式的解析差异"收口在 parser 内部，
# 让后续 ingest/index 链只面对统一的 `Chunk` 契约。


class UnsupportedFileTypeError(ValueError):
    """不支持的文件格式异常。"""

    def __init__(self, ext: str, supported: List[str]):
        # 报错时直接把当前扩展和支持列表暴露出来，方便前端提示。
        message = f"不支持的文件格式: {ext or '(none)'}。支持: {', '.join(supported)}"
        super().__init__(message)


class StructuredChunkingError(ValueError):
    """Strict structured-ingest failure with stable upload-visible fields."""

    def __init__(
        self,
        code: str,
        *,
        paper_id: str | None = None,
        file_id: str | None = None,
        detail: str = "",
    ) -> None:
        self.stage: Literal["parse"] = "parse"
        self.code = code
        self.paper_id = paper_id or None
        self.file_id = file_id or None
        self.detail = detail
        super().__init__(self._message())

    def _message(self) -> str:
        parts = [f"stage={self.stage}", f"code={self.code}"]
        if self.paper_id:
            parts.append(f"paper_id={self.paper_id}")
        if self.file_id:
            parts.append(f"file_id={self.file_id}")
        if self.detail:
            parts.append(f"detail={self.detail}")
        return " ".join(parts)


class DocumentParser(ABC):
    """文档解析器抽象基类，约束所有格式统一输出 `List[Chunk]`。"""

    def parse(self, file_path: str, paper_id: str) -> List[Chunk]:
        """参数为文件路径、论文身份；返回分块列表，完整来源由 parse_document 一起交付。"""
        return self.parse_document(file_path, paper_id)["chunks"]

    def parse_document(self, file_path: str, paper_id: str) -> dict[str, Any]:
        """参数为文件路径、论文身份；返回含规范来源项目和内容版本的解析文档。"""
        mode = Path(file_path).suffix.lower().lstrip(".") + "_parser_local"
        return _sections_to_document(self.collect_sections(file_path), paper_id, parser_mode=mode)

    @abstractmethod
    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为文件路径；返回保留原始解析位置的章节片段。"""

    @abstractmethod
    def supported_extensions(self) -> List[str]:
        """返回支持的文件扩展名。"""


class ParserRegistry:
    """按扩展名路由解析器，是当前多格式上传链的解析入口。"""

    def __init__(self) -> None:
        # key 是 `.pdf` 这种扩展名，value 是对应解析器实例。
        self._parsers: dict[str, DocumentParser] = {}

    def register(self, parser: DocumentParser) -> None:
        """注册一个解析器，并绑定其支持的全部扩展名。"""
        for ext in parser.supported_extensions():
            # 扩展名统一转小写，避免大小写导致路由失败。
            self._parsers[ext.lower()] = parser

    def parse(self, file_path: str, paper_id: str) -> List[Chunk]:
        """根据文件扩展名选择解析器并执行统一解析。"""
        ext = Path(file_path).suffix.lower()
        parser = self._parsers.get(ext)
        if parser is None:
            raise UnsupportedFileTypeError(ext, self.supported_types())
        # 各解析器最终都必须返回统一的 `List[Chunk]`。
        return parser.parse(file_path, paper_id)

    def supported_types(self) -> List[str]:
        """返回当前注册表支持的全部扩展名。"""
        return sorted(self._parsers.keys())

    def parse_document(self, file_path: str, paper_id: str) -> dict[str, Any]:
        """参数为文件路径与论文身份；选择解析器并返回正文、来源和内容版本。"""
        ext = Path(file_path).suffix.lower()
        parser = self._parsers.get(ext)
        if parser is None:
            raise UnsupportedFileTypeError(ext, self.supported_types())
        return parser.parse_document(file_path, paper_id)


class PdfParser(DocumentParser):
    """PDF 解析器，当前上传主链对 PDF 文件优先走这里。"""

    def supported_extensions(self) -> List[str]:
        return [".pdf"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 PDF 路径；返回逐页提取的正文片段，文档句柄在完成后关闭。"""
        from app.services.pdf_document import open_pdf

        # 1. 【文档解析】【PDF 分块】先抽取真实页，再交给现有父子块和来源定位流程
        with open_pdf(file_path) as document:
            return _collect_pdf_sections(document)


class DocxParser(DocumentParser):
    """DOCX 解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 DOCX 路径；返回标题、正文和表格片段。"""
        from docx import Document

        # DOCX 先按标题段落拆 section，再交给统一 chunk 流程。
        doc = Document(file_path)
        sections = _collect_docx_sections(doc)
        return sections


class PptxParser(DocumentParser):
    """PPTX 解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".pptx"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 PPTX 路径；返回带实际幻灯片页码的正文片段。"""
        from pptx import Presentation

        # PPTX 按 slide 抽文本，每页 slide 可以对应一个 section。
        presentation = Presentation(file_path)
        sections = _collect_pptx_sections(presentation)
        return sections


class MarkdownParser(DocumentParser):
    """Markdown 解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".md", ".markdown"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 Markdown 路径；返回带原文字符位置的结构片段。"""
        # Markdown 统一走 block-aware 分块，尽量保留 heading/list/table/formula/code 边界。
        return _collect_markdown_sections(_read_text_file(file_path))


class HtmlParser(DocumentParser):
    """HTML 解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".html", ".htm"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 HTML 路径；返回去除导航噪音后的正文结构。"""
        from bs4 import BeautifulSoup

        # 先清除 script/style/nav 等噪音节点，再抽正文结构。
        soup = BeautifulSoup(_read_text_file(file_path), "html.parser")
        _strip_html_noise(soup)
        sections = _collect_html_sections(soup)
        return sections


class TxtParser(DocumentParser):
    """纯文本解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".txt", ".text"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为纯文本路径；返回单一正文片段，不编造物理页码。"""
        text = _read_text_file(file_path).strip()
        if not text:
            return []
        # 纯文本没有天然结构，整体作为一个 section 进入后续切块。
        return [_section("content", text, 0)]


class LatexParser(DocumentParser):
    """LaTeX 源文件解析器。"""

    def supported_extensions(self) -> List[str]:
        return [".tex", ".latex"]

    def collect_sections(self, file_path: str) -> List[dict]:
        """参数为 LaTeX 路径；返回按 section 标记提取的正文片段。"""
        # 先剥掉注释，再按 section/subsection 结构切分。
        text = _strip_tex_comments(_read_text_file(file_path))
        sections = _split_latex_sections(text)
        return sections


def get_parser_registry() -> ParserRegistry:
    """返回默认解析注册表单例，供上传入库主链复用。"""
    global _REGISTRY
    if _REGISTRY is None:
        # 首次访问时懒加载全部默认解析器。
        registry = ParserRegistry()
        for parser in _build_default_parsers():
            registry.register(parser)
        _REGISTRY = registry
    return _REGISTRY


def _build_default_parsers() -> List[DocumentParser]:
    """构建默认启用的解析器集合。"""
    # 当前默认支持的多格式上传集合都在这里集中注册。
    return [
        PdfParser(),
        DocxParser(),
        PptxParser(),
        MarkdownParser(),
        HtmlParser(),
        TxtParser(),
        LatexParser(),
    ]


def _collect_docx_sections(doc) -> List[dict]:
    """按标题和正文段落抽取 DOCX sections。"""
    sections: List[dict] = []
    current_title = "untitled"
    current_lines: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = getattr(getattr(para, "style", None), "name", "") or ""
        # Heading 段落视作新的 section 起点。
        if style_name.startswith("Heading"):
            sections = _flush_section(sections, current_title, current_lines, 0)
            current_title = text
            current_lines = []
            continue
        current_lines.append(text)
    sections = _flush_section(sections, current_title, current_lines, 0)
    for table in doc.tables:
        # 表格转换成 Markdown，作为额外 section 追加。
        table_text = _table_to_markdown(table.rows)
        if table_text:
            sections.append(_section("table", table_text, 0, block_type="table"))
    return sections


def _collect_pptx_sections(presentation) -> List[dict]:
    """把 PPTX 每张 slide 转成 section。"""
    sections: List[dict] = []
    for page_num, slide in enumerate(presentation.slides, start=1):
        title, body_lines = _extract_slide_text(slide)
        if body_lines or title:
            text = "\n".join(body_lines) if body_lines else title
            sections.append(_section(title or f"Slide {page_num}", text, page_num))
    return sections


def _extract_slide_text(slide) -> tuple[str, List[str]]:
    """提取单张 slide 的标题、正文和表格文本。"""
    title = ""
    body_lines: List[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if not text:
                continue
            # 第一块标题占位符文本优先视作 slide 标题。
            if _is_slide_title(shape, title):
                title = text
            else:
                body_lines.append(text)
        if getattr(shape, "has_table", False):
            table_text = _table_to_markdown(shape.table.rows)
            if table_text:
                body_lines.append(table_text)
    return title, body_lines


def _is_slide_title(shape, current_title: str) -> bool:
    """判断当前 shape 是否可视为 slide 标题。"""
    if current_title:
        return False
    placeholder = getattr(shape, "placeholder_format", None)
    return bool(placeholder and getattr(placeholder, "type", None) == 1)


def _split_markdown_sections(text: str) -> List[dict]:
    """按 Markdown 标题切分 sections。"""
    header_pattern = re.compile(r"^(#{1,3})\s+(.+)$", re.MULTILINE)
    return _split_structured_sections(text, header_pattern, lambda match: match.group(2).strip())


def parse_markdown_text(
    text: str,
    paper_id: str,
    *,
    content_list: list[dict[str, Any]] | None = None,
    parser_mode: str = "",
) -> List[Chunk]:
    """把 Markdown 文本解析成 block-aware 的 Chunk 列表。"""
    return parse_markdown_document(text, paper_id, content_list=content_list, parser_mode=parser_mode)["chunks"]


def parse_markdown_document(
    text: str, paper_id: str, *, content_list: list[dict[str, Any]] | None = None, parser_mode: str = "",
    text_normalizer: Callable[[str], str] | None = None,
) -> dict[str, Any]:
    """参数为原始 Markdown、论文和可选块清洗器；返回正文/来源，原始字符定位在清洗前固定。"""
    sections = _collect_markdown_sections(text, content_list=content_list, text_normalizer=text_normalizer)
    return _sections_to_document(sections, paper_id, parser_mode=parser_mode)


def _collect_markdown_sections(
    text: str,
    *,
    content_list: list[dict[str, Any]] | None = None,
    text_normalizer: Callable[[str], str] | None = None,
) -> List[dict]:
    if content_list:
        return _collect_markdown_sections_from_content_list(content_list)
    return _collect_markdown_sections_from_text(text, text_normalizer=text_normalizer)


def _collect_markdown_sections_from_content_list(content_list: list[dict[str, Any]]) -> List[dict]:
    """参数为规范内容列表；返回逐项目保留来源的章节，不把整节归到首块页码。"""
    blocks = []
    for index, item in enumerate(content_list):
        block = _content_item_to_markdown_block(item)
        if block is not None:
            blocks.append(_source_block(block, source_kind="content_list", source_index=index))
    return _blocks_to_sections(blocks)


def _collect_markdown_sections_from_text(text: str, *, text_normalizer: Callable[[str], str] | None = None) -> List[dict]:
    """参数为 Markdown 正文；返回带原始字符区间的章节，无页面标记时不提供物理页码。"""
    lines = text.splitlines(keepends=True)
    line_offsets = [0]
    for line in lines:
        line_offsets.append(line_offsets[-1] + len(line))
    lines = [line.rstrip("\r\n") for line in lines]
    blocks = []
    page_num = 0
    index = 0
    while index < len(lines):
        raw_line = lines[index].rstrip("\r\n")
        stripped = raw_line.strip()
        if not stripped:
            index += 1
            continue
        page_marker = _MARKDOWN_PAGE_MARKER.match(stripped)
        if page_marker:
            page_num = int(page_marker.group(1))
            index += 1
            continue
        heading = _MARKDOWN_HEADING_PATTERN.match(raw_line)
        if heading:
            title = text_normalizer(heading.group(2)) if text_normalizer else heading.group(2)
            blocks.append({"type": "heading", "text": _clean_inline_markdown_text(title),
                "level": len(heading.group(1))})
            index += 1
            continue
        body, next_index = _consume_markdown_block(lines, index)
        if text_normalizer is not None:
            body = text_normalizer(body)
        if body:
            block = {"type": _infer_markdown_block_type(body), "text": body, "page_num": page_num}
            sourced = _source_block(block, source_kind="markdown", source_index=index)
            source_item = next(iter(sourced["source_items"].values()))
            source_item["source_start"] = line_offsets[index]
            source_item["source_end"] = line_offsets[next_index]
            blocks.append(sourced)
        index = next_index
    return _blocks_to_sections(blocks)


def _source_block(block: dict, *, source_kind: str, source_index: int) -> dict:
    """参数为转换后的块与原始项目位置；返回规范阅读正文、项目身份及起始来源映射。"""
    text = str(block["text"]).strip()
    source_id = f"src-v1:{source_kind}:{source_index}"
    page_num = int(block.get("page_num") or 0) or None
    item = {"source_item": source_id, "text": text, "source_kind": source_kind,
        "source_index": source_index, "block_type": block["type"], "page_num": page_num}
    fragment: SourceFragment = {"source_item": source_id, "source_start": 0, "source_end": len(text),
        "text_start": 0, "text_end": len(text), "page_num": page_num}
    return {**block, "text": text, "source_items": {source_id: item}, "source_fragments": [fragment]}


def _blocks_to_sections(blocks: list[dict]) -> List[dict]:
    """参数为有序解析块；返回章节片段，标签随表格/公式一起分组并保留各自来源。"""
    sections = []
    prose: list[dict] = []
    heading_stack: list[str] = []
    title = "untitled"
    for block in blocks:
        if block["type"] == "heading":
            if prose and not _section_is_excluded(title):
                sections.append(_blocks_section(prose, title, "paragraph"))
            prose = []
            level = max(int(block.get("level") or 1), 1)
            heading_stack = heading_stack[:level - 1] + [block["text"]]
            title = " / ".join(heading_stack)
            continue
        if _section_is_excluded(title):
            continue
        block_type = _normalized_block_type(block["type"])
        if block_type not in _ATOMIC_BLOCK_TYPES:
            prose.append(block)
            continue
        # 1. 【文档解析】【结构分组】相邻标签属于原子结构，移动标签时同步移动其来源
        attached = []
        if prose and _is_atomic_label(prose[-1]["text"], block_type):
            attached = [prose[-1]]
            prose = prose[:-1]
        if prose:
            sections.append(_blocks_section(prose, title, "paragraph"))
            prose = []
        sections.append(_blocks_section([*attached, block], title, block_type))
    if prose and not _section_is_excluded(title):
        sections.append(_blocks_section(prose, title, "paragraph"))
    return sections


def _blocks_section(blocks: list[dict], title: str, block_type: str) -> dict:
    """参数为同组正文块、章节和类型；返回正文合并结果及逐项目区间，分隔符不映射来源。"""
    value = _join_mapped((_MappedText(block["text"], tuple(block["source_fragments"]))
        for block in blocks), _PARAGRAPH_SEPARATOR)
    items = {key: item for block in blocks for key, item in block["source_items"].items()}
    pages = {item["page_num"] for item in items.values()}
    page_num = next(iter(pages)) if len(pages) == 1 else 0
    section = _section(title, value.text, page_num or 0, block_type=block_type,
        section_path=title, asset_ref=str(blocks[-1].get("asset_ref") or ""))
    return {**section, "source_items": items, "source_fragments": list(value.fragments)}


def _is_atomic_label(text: str, block_type: str) -> bool:
    """参数为前一段正文和原子类型；返回该段是否为对应公式或表格的关联标签。"""
    return (block_type == "formula" and _looks_like_formula_label(text)) or (
        block_type == "table" and _looks_like_table_caption(text))


def _looks_like_formula_label(text: str) -> bool:
    """参数为段落正文；返回是否带有可关联显示公式的指标标签。"""
    return bool(_FORMULA_LABEL_PATTERN.search(str(text or "").strip()))


def _looks_like_table_caption(text: str) -> bool:
    """参数为段落正文；返回是否为带编号的表格标题。"""
    return bool(_TABLE_CAPTION_PATTERN.match(str(text or "").strip()))


def _normalized_block_type(block_type: str) -> str:
    type_map = {"paragraph": "paragraph", "list": "paragraph", "code": "code",
                "table": "table", "formula": "formula", "figure_caption": "figure_caption",
                "heading": "heading", "caption": "figure_caption"}
    return type_map.get(block_type, block_type or "paragraph")


def _infer_markdown_block_type(block: str) -> str:
    stripped = str(block or "").strip()
    if stripped.startswith(("```", "~~~")):
        return "code"
    if stripped.startswith("$$") and stripped.endswith("$$"):
        return "formula"
    if stripped.startswith("![") or stripped.lower().startswith("figure "):
        return "figure_caption"
    return "table" if _looks_like_markdown_table_line(stripped.splitlines()[0]) else "paragraph"


def _dominant_block_type(block_types: list[str]) -> str:
    if not block_types:
        return "paragraph"
    type_map = {"paragraph": "paragraph", "list": "paragraph", "code": "paragraph",
                "table": "table", "formula": "formula", "figure_caption": "figure_caption",
                "heading": "heading", "caption": "figure_caption"}
    for bt in block_types:
        mapped = type_map.get(bt, bt)
        if mapped != "paragraph":
            return mapped
    return "paragraph"


def _consume_markdown_block(lines: list[str], start_index: int) -> tuple[str, int]:
    """参数为原始行和起始行号；返回一个完整结构块及下一位置，不越过页面或章节标记。"""
    stripped = lines[start_index].strip()
    if stripped.startswith(("```", "~~~")) or stripped == "$$":
        fence = "$$" if stripped == "$$" else stripped[:3]
        return _consume_fenced_block(lines, start_index, fence)
    if _looks_like_markdown_table_line(stripped):
        block_lines = [lines[start_index].rstrip()]
        index = start_index + 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or not _looks_like_markdown_table_line(candidate):
                break
            block_lines.append(lines[index].rstrip())
            index += 1
        return "\n".join(block_lines).strip(), index
    if _looks_like_markdown_list_item(stripped):
        block_lines = [lines[start_index].rstrip()]
        index = start_index + 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                index += 1
                break
            if _looks_like_markdown_list_item(candidate) or candidate.startswith("  ") or candidate.startswith("\t"):
                block_lines.append(lines[index].rstrip())
                index += 1
                continue
            break
        return "\n".join(block_lines).strip(), index
    return _consume_markdown_paragraph(lines, start_index)


def _consume_fenced_block(lines: list[str], start_index: int, fence: str) -> tuple[str, int]:
    """参数为行列表、起点和围栏；返回保持完整的代码/公式及结束行位置。"""
    block_lines = [lines[start_index]]
    index = start_index + 1
    while index < len(lines):
        block_lines.append(lines[index])
        candidate = lines[index].strip()
        index += 1
        closed = candidate == fence if fence == "$$" else candidate.startswith(fence)
        if closed:
            break
    return "\n".join(block_lines).strip(), index


def _consume_markdown_paragraph(lines: list[str], start_index: int) -> tuple[str, int]:
    """参数为行列表和段落起点；返回归一段落及下一位置，保留原始位置供来源项目定位。"""
    paragraph_lines = [lines[start_index].rstrip()]
    index = start_index + 1
    while index < len(lines):
        candidate = lines[index]
        candidate_stripped = candidate.strip()
        if not candidate_stripped:
            index += 1
            break
        if _MARKDOWN_HEADING_PATTERN.match(candidate) or _MARKDOWN_PAGE_MARKER.match(candidate_stripped):
            break
        if candidate_stripped.startswith(("```", "~~~")) or candidate_stripped == "$$":
            break
        if _looks_like_markdown_table_line(candidate_stripped) or _looks_like_markdown_list_item(candidate_stripped):
            break
        paragraph_lines.append(candidate.rstrip())
        index += 1
    normalized = _clean_inline_markdown_text(" ".join(line.strip() for line in paragraph_lines if line.strip()))
    return normalized, index


def _looks_like_markdown_list_item(line: str) -> bool:
    return bool(_LIST_ITEM_PATTERN.match(line))


def _looks_like_markdown_table_line(line: str) -> bool:
    if not _TABLE_ROW_PATTERN.match(line):
        return False
    parts = [part.strip() for part in line.split("|")[1:-1]]
    return bool(parts)


def _section_is_excluded(title: str) -> bool:
    return bool(_EXCLUDED_SECTION_PATTERN.search(str(title or "")))


def _content_item_to_markdown_block(item: dict[str, Any]) -> dict[str, Any] | None:
    """参数为原始内容项目；返回规范正文块，噪音项目不进入知识正文。"""
    block_type = str(item.get("type") or item.get("block_type") or "").strip().lower()
    if block_type in _MINERU_NOISE_TYPES:
        return None
    page_num = _coerce_page_num(item.get("page_idx"))
    if block_type in {"title", "heading"}:
        text = _clean_inline_markdown_text(_stringify_markdown_value(item.get("text") or item.get("title")))
        if not text:
            return None
        return {"type": "heading", "text": text, "level": 1, "page_num": page_num}
    if block_type == "text":
        text = _clean_inline_markdown_text(_stringify_markdown_value(item.get("text")))
        if not text:
            return None
        level = _coerce_int(item.get("text_level"))
        if level and level > 0:
            return {"type": "heading", "text": text, "level": level, "page_num": page_num}
        return {"type": "paragraph", "text": text, "page_num": page_num}
    if block_type == "list":
        text = _clean_inline_markdown_text(_stringify_markdown_value(item.get("text") or item.get("list_body")))
        if not text:
            return None
        return {"type": "list", "text": text, "page_num": page_num}
    if block_type in {"table", "equation", "formula", "code", "image", "figure", "chart"}:
        return _structured_content_item(item, block_type, page_num)
    text = _clean_inline_markdown_text(_stringify_markdown_value(item))
    return {"type": block_type or "paragraph", "text": text, "page_num": page_num} if text else None


def _structured_content_item(item: dict[str, Any], block_type: str, page_num: int) -> dict[str, Any] | None:
    """参数为原子项目、类型与实际页码；返回保留表格/公式/代码/图注正文的统一结构块。"""
    if block_type == "table":
        caption = _clean_inline_markdown_text(_stringify_markdown_value(item.get("table_caption") or item.get("caption")))
        raw_html = _stringify_markdown_value(item.get("table_body") or item.get("html"))
        footnote = _clean_inline_markdown_text(_stringify_markdown_value(item.get("table_footnote") or item.get("footnote")))
        text = "\n".join(part for part in [caption, _html_table_rows(raw_html), footnote] if part).strip()
        if not text:
            return None
        return {"type": "table", "text": text, "page_num": page_num}
    if block_type in {"equation", "formula"}:
        text = _stringify_markdown_value(item.get("latex") or item.get("text") or item.get("formula"))
        text = text.strip()
        if not text:
            return None
        if not text.startswith("$$"):
            text = f"$$\n{text}\n$$"
        return {"type": "formula", "text": text, "page_num": page_num}
    if block_type == "code":
        caption = _clean_inline_markdown_text(_stringify_markdown_value(item.get("code_caption") or item.get("caption")))
        body = _stringify_markdown_value(item.get("code_body") or item.get("text"))
        body = body.strip()
        if not body:
            return None
        fenced = body if body.startswith("```") else f"```\n{body}\n```"
        text = "\n\n".join(part for part in [caption, fenced] if part).strip()
        return {"type": "code", "text": text, "page_num": page_num}
    if block_type in {"image", "figure", "chart"}:
        caption = _clean_inline_markdown_text(
            _stringify_markdown_value(item.get("img_caption") or item.get("figure_caption") or item.get("caption"))
        )
        nearby = _clean_inline_markdown_text(
            _stringify_markdown_value(item.get("img_footnote") or item.get("near_text") or item.get("footnote"))
        )
        if block_type == "chart":
            # 1. 【文档解析】【图题保真】原生图题和脚注各自优先，空内容仍沿用既有别名
            caption = _clean_inline_markdown_text(_stringify_markdown_value(item.get("chart_caption"))) or caption
            nearby = _clean_inline_markdown_text(_stringify_markdown_value(item.get("chart_footnote"))) or nearby
        text = "\n".join(part for part in [caption, nearby] if part).strip()
        if not text:
            return None
        asset_ref = str(item.get("img_path") or item.get("image_path") or item.get("path") or "").strip()
        return {"type": "figure_caption", "text": text, "page_num": page_num, "asset_ref": asset_ref}
    raise ValueError(f"不支持的结构块类型: {block_type}")


def _coerce_int(value: Any) -> int | None:
    if isinstance(value, int):
        return value
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _coerce_page_num(value: Any) -> int:
    """参数为原始零起始页索引；返回实际的一起始页码，缺省为零，非法索引明确失败。"""
    if value is None:
        return 0
    page_idx = _coerce_int(value)
    if type(page_idx) is not int or page_idx < 0:
        raise ValueError("来源页索引必须是非负整数")
    return page_idx + 1


def _stringify_markdown_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, list):
        return "\n".join(_stringify_markdown_value(item) for item in value if _stringify_markdown_value(item).strip())
    if isinstance(value, dict):
        preferred_keys = [
            "text",
            "caption",
            "title",
            "content",
            "html",
            "latex",
            "code_body",
        ]
        parts = [
            _stringify_markdown_value(value.get(key))
            for key in preferred_keys
            if _stringify_markdown_value(value.get(key)).strip()
        ]
        if parts:
            return "\n".join(parts)
    return ""


def _clean_inline_markdown_text(text: str) -> str:
    normalized = str(text or "").replace("\r\n", "\n").strip()
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _html_table_rows(html: str) -> str:
    """
    按行还原 HTML 表格结构，作为规范阅读正文。

    参数 html 为 MinerU 原始 table_body；返回每行一行的文本，行内单元格用 ` | ` 分隔。
    合并单元格只标注跨度，不复制内容，避免把同一数值重复成多列。
    """
    if not html.strip():
        raise ValueError("表格正文缺少可读取的行列结构")
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    lines: list[str] = []
    for row in soup.find_all("tr"):
        cells = row.find_all(_TABLE_CELL_TAGS)
        if not cells:
            continue
        # 1. 单元格跨度原样保留，列对应关系不因合并而错位
        rendered = [_html_cell_text(cell) for cell in cells]
        lines.append("| " + " | ".join(rendered) + " |")
    if not lines:
        raise ValueError("表格正文未包含可读取的行列结构")
    return "\n".join(lines)


def _html_cell_text(cell) -> str:
    """参数为单个单元格；返回去掉换行后的文本，并在有合并时追加跨度标注。"""
    text = re.sub(r"\s+", " ", cell.get_text(" ", strip=True)).strip()
    spans = [f"{name}={cell.get(name)}" for name in ("rowspan", "colspan") if cell.get(name)]
    if not spans:
        return text
    return f"{text} [{' '.join(spans)}]".strip()


def _split_latex_sections(text: str) -> List[dict]:
    """按 LaTeX section/subsection 命令切分 sections。"""
    pattern = re.compile(r"\\(?:section|subsection|subsubsection)\{([^}]+)\}")
    sections = _split_structured_sections(text, pattern, lambda match: match.group(1).strip())
    return [_section(item["title"], _clean_latex(item["text"]), item["page_num"]) for item in sections if _clean_latex(item["text"])]


def _split_structured_sections(text: str, pattern, title_getter) -> List[dict]:
    """按结构化标题模式把长文本切成 section 列表。"""
    sections: List[dict] = []
    last_end = 0
    last_title = "untitled"
    for match in pattern.finditer(text):
        # 当前标题之前的内容归到上一个 section。
        content = text[last_end:match.start()].strip()
        if content:
            sections.append(_section(last_title, content, 0))
        last_title = title_getter(match)
        last_end = match.end()
    tail = text[last_end:].strip()
    if tail:
        sections.append(_section(last_title, tail, 0))
    return sections


def _strip_html_noise(soup) -> None:
    """移除 HTML 中明显非正文的噪音标签。"""
    for tag in soup.find_all(["script", "style", "nav", "footer", "header"]):
        tag.decompose()


def _collect_html_sections(soup) -> List[dict]:
    """按 HTML 标题和文本节点提取 sections。"""
    sections: List[dict] = []
    current_title = "untitled"
    current_lines: List[str] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote"]):
        text = element.get_text(separator=" ", strip=True)
        if not text:
            continue
        # 标题标签开始新的 section，其他标签累计进当前正文。
        if element.name.startswith("h"):
            sections = _flush_section(sections, current_title, current_lines, 0)
            current_title = text
            current_lines = []
            continue
        current_lines.append(text)
    return _flush_section(sections, current_title, current_lines, 0)


def _flush_section(sections: List[dict], title: str, lines: List[str], page_num: int) -> List[dict]:
    """把缓存中的 section 内容刷入列表。"""
    text = "\n".join(lines).strip()
    if text:
        sections.append(_section(title, text, page_num))
    return sections


def _section(
    title: str,
    text: str,
    page_num: int,
    *,
    block_type: str = "paragraph",
    section_path: str = "",
    asset_ref: str = "",
    embedding_text: str = "",
) -> dict:
    return {
        "title": title or "untitled",
        "text": text,
        "page_num": page_num,
        "block_type": block_type,
        "section_path": section_path or title or "",
        "asset_ref": asset_ref,
        # 阅读正文与索引表达分离，索引表达不能替代规范来源文本
        "embedding_text": embedding_text,
    }


class _HFTokenCounter:
    """基于 HuggingFace tokenizer 的轻量 token 计数器。"""

    def __init__(self, tokenizer) -> None:
        """参数为本地模型已有分词器；保存实际计数规则和版本身份。"""
        self._tokenizer = tokenizer
        self._sha256 = hashlib.sha256(tokenizer.backend_tokenizer.to_str().encode("utf-8")).hexdigest()

    def count(self, text: str) -> int:
        """参数为原文；保留本地模型的既有无特殊标记计数方式。"""
        return len(self._tokenizer.encode(text, add_special_tokens=False))

    def identity(self) -> dict[str, str | bool]:
        """无参数；返回本地模型实际分词规则身份，不用模型别名代替文件内容。"""
        return {"engine": f"transformers:{version('transformers')}", "sha256": self._sha256,
                "add_special_tokens": False}



@lru_cache(maxsize=1)
def _get_token_counter() -> _HFTokenCounter | ApiTokenCounter:
    """无参数；按明确模型模式加载一次计数器，API 仅使用独立分词文件。"""
    if settings.embedding_mode == "api":
        return load_api_token_counter(settings.chunk_tokenizer_path, settings.chunk_tokenizer_sha256)
    try:
        from transformers import AutoTokenizer
    except ImportError as exc:
        raise RuntimeError("缺少 transformers 依赖，无法执行 token-aware 分块") from exc
    configure_offline_env()
    try:
        tokenizer = AutoTokenizer.from_pretrained(
            resolve_local_model_source(settings.embedding_model),
            local_files_only=True,
            use_fast=True,
        )
    except Exception as exc:
        raise RuntimeError(
            f"无法从本地缓存加载 tokenizer: {settings.embedding_model}"
        ) from exc
    return _HFTokenCounter(tokenizer)


def _count_tokens(text: str) -> int:
    """返回文本 token 数。"""
    if not text.strip():
        return 0
    return _get_token_counter().count(text)


def _token_budget(max_tokens: int, overlap_tokens: int) -> int:
    """计算扣除 overlap 后的有效内容预算。"""
    return max(1, max_tokens - max(0, overlap_tokens))


def _overlap_mapped(parts: list[_MappedText], overlap_tokens: int, max_tokens: int, separator: str) -> list[_MappedText]:
    """参数为连续正文、重叠与总额度和分隔符；返回带来源的重叠正文，永不裁掉当前新增材料。"""
    if overlap_tokens <= 0 or len(parts) < 2:
        return parts
    merged = [parts[0]]
    for previous, current in zip(parts, parts[1:]):
        # 1. 【文档解析】【重叠分块】只在原正文字符坐标上选后缀，不经 token 解码改变原文
        left, right = 0, len(previous.text)
        while left < right:
            middle = (left + right) // 2
            suffix = previous.slice(middle, len(previous.text)).strip()
            candidate = _join_mapped([suffix, current], separator)
            fits = _count_tokens(suffix.text) <= overlap_tokens and _count_tokens(candidate.text) <= max_tokens
            if fits:
                right = middle
            else:
                left = middle + 1
        overlap = previous.slice(left, len(previous.text)).strip()
        merged.append(_join_mapped([overlap, current], separator))
    return merged


def _sections_to_chunks(sections: List[dict], paper_id: str, *, parser_mode: str = "") -> List[Chunk]:
    """参数为章节、论文和解析模式；返回携带逐字符来源映射的父子分块。"""
    return _sections_to_document(sections, paper_id, parser_mode=parser_mode)["chunks"]


def _sections_to_document(sections: List[dict], paper_id: str, *, parser_mode: str = "") -> dict[str, Any]:
    """参数为章节、论文和解析模式；返回全部规范来源及分块，不在分块后反查原文。"""
    sourced = [_ensure_section_sources(section, index, parser_mode) for index, section in enumerate(sections)]
    items = {key: item for section in sourced for key, item in section["source_items"].items()}
    merged = _merge_sections_for_parenting(sourced, paper_id)
    chunks: List[Chunk] = []
    parent_idx = 0
    for section in merged:
        for parent in _parent_values_for_section(section, paper_id):
            if not parent.text:
                continue
            chunks.extend(_split_to_child_chunks(
                parent_text=parent.text, parent_id=f"{paper_id}_{parent_idx}",
                page_num=section["page_num"], section=section["title"],
                target_size=settings.child_chunk_target_size,
                overlap_tokens=settings.child_chunk_overlap_tokens,
                block_type=section.get("block_type", "paragraph"),
                section_path=section.get("section_path", ""), asset_ref=section.get("asset_ref", ""),
                parser_mode=parser_mode, paper_id=paper_id,
                source_embedding_text=section.get("embedding_text", ""),
                parent_source_fragments=parent.fragments))
            parent_idx += 1
    return {"chunks": chunks, "source_items": items, "content_schema_version": CONTENT_SCHEMA_VERSION}


def _ensure_section_sources(section: dict, index: int, parser_mode: str) -> dict:
    """参数为未分块章节、顺序和解析模式；为本地格式固定来源项目，已有逐块来源保持原样。"""
    if "source_items" in section:
        return dict(section)
    block = {"text": section["text"], "type": section.get("block_type", "paragraph"),
        "page_num": section.get("page_num")}
    sourced = _source_block(block, source_kind=parser_mode or "local_section", source_index=index)
    return {**section, "text": sourced["text"], "source_items": sourced["source_items"],
        "source_fragments": sourced["source_fragments"]}


def _merge_sections_for_parenting(sections: List[dict], paper_id: str) -> List[dict]:
    clean_sections = [dict(item) for item in sections if str(item.get("text") or "").strip()]
    _raise_for_invalid_structure(clean_sections, paper_id)
    if not getattr(settings, "structure_merge_enabled", True):
        return clean_sections
    return _merge_structure_sections(clean_sections, settings.chunk_merge_floor)


def _parent_texts_for_section(section: dict, paper_id: str) -> List[str]:
    """参数为章节和论文身份；返回父块正文，离线报告与生产分块共用同一切分逻辑。"""
    return [value.text for value in _parent_values_for_section(section, paper_id)]


def _parent_values_for_section(section: dict, paper_id: str) -> list[_MappedText]:
    """参数为带来源章节和论文身份；返回保留父级重叠范围的正文片段。"""
    value = _MappedText(str(section.get("text") or ""), tuple(section.get("source_fragments", ()))).strip()
    if _is_atomic_section(section):
        _raise_if_atomic_block_too_large(section, paper_id)
        return [value]
    return _split_mapped_section(value, max_size=settings.parent_chunk_max_size,
        overlap_tokens=settings.parent_chunk_overlap_tokens, paper_id=paper_id, section_path=_section_path(section))


def _raise_for_invalid_structure(sections: list[dict], paper_id: str) -> None:
    if not sections:
        return
    keys = [_section_structure_key(section) for section in sections]
    usable_keys = [key for key in keys if _is_usable_structure_key(key)]
    if _has_single_unusable_section_key(keys):
        detail = f"single unusable section key: {keys[0]}"
        _raise_chunking_error(_ERROR_SINGLE_SECTION_KEY, paper_id, detail)
    if not usable_keys:
        _raise_chunking_error(_ERROR_MISSING_SECTION_PATH, paper_id, "no usable section_path")


def _has_single_unusable_section_key(keys: list[str]) -> bool:
    return len(keys) > 1 and len(set(keys)) == 1 and not _is_usable_structure_key(keys[0])


def _is_usable_structure_key(key: str) -> bool:
    if not key or key == "top:untitled":
        return False
    if key.startswith("top:"):
        label = key.removeprefix("top:").strip()
        return label not in _GENERIC_SECTION_PATHS
    return True


def _merge_structure_sections(sections: List[dict], floor: int) -> List[dict]:
    merged: list[dict] = []
    buffer: Optional[dict] = None
    for section in sections:
        if _is_atomic_section(section):
            if buffer is not None:
                merged.append(buffer)
                buffer = None
            current = dict(section)
            if _can_merge_atomic_neighbors(merged, current):
                merged[-1] = _combine_sections(merged[-1], current)
            else:
                merged.append(current)
            continue
        buffer = _next_structure_buffer(buffer, section, floor, merged)
    if buffer is not None:
        merged.append(buffer)
    return _merge_trailing_structure_stubs(merged, floor)


def _next_structure_buffer(buffer: Optional[dict], section: dict, floor: int, merged: list[dict]) -> dict:
    current = dict(section)
    if buffer is None:
        return current
    if not _can_merge_structure_neighbors(buffer, current):
        merged.append(buffer)
        return current
    if _count_tokens(str(buffer.get("text") or "")) < floor:
        return _combine_sections(buffer, current)
    merged.append(buffer)
    return current


def _merge_trailing_structure_stubs(sections: list[dict], floor: int) -> list[dict]:
    merged: list[dict] = []
    for section in sections:
        if _can_merge_backward(section, merged, floor):
            merged[-1] = _combine_sections(merged[-1], section)
            continue
        merged.append(section)
    return merged


def _can_merge_backward(section: dict, merged: list[dict], floor: int) -> bool:
    if not merged or _is_atomic_section(section) or _is_atomic_section(merged[-1]):
        return False
    if _count_tokens(str(section.get("text") or "")) >= floor:
        return False
    return _can_merge_structure_neighbors(merged[-1], section)


def _can_merge_atomic_neighbors(merged: list[dict], current: dict) -> bool:
    if not merged or not _is_atomic_section(merged[-1]):
        return False
    if str(merged[-1].get("block_type") or "") != "formula":
        return False
    if str(current.get("block_type") or "") != "formula":
        return False
    if _section_path(merged[-1]) != _section_path(current):
        return False
    combined_text = f'{merged[-1]["text"]}{_PARAGRAPH_SEPARATOR}{current["text"]}'.strip()
    return _count_tokens(combined_text) <= settings.parent_chunk_max_size


def _combine_sections(left: dict, right: dict) -> dict:
    """参数为可合并的相邻章节；返回新章节，正文、来源项目和区间同时合并。"""
    value = _join_mapped((_MappedText(item["text"], tuple(item.get("source_fragments", ())))
        for item in (left, right)), _PARAGRAPH_SEPARATOR)
    combined = {**left, "text": value.text,
        "title": _combine_section_label(str(left.get("title") or ""), str(right.get("title") or "")),
        "section_path": _combine_section_label(_section_path(left), _section_path(right))}
    if "source_items" in left:
        combined["source_items"] = {**left["source_items"], **right["source_items"]}
        combined["source_fragments"] = list(value.fragments)
    combined.pop("embedding_text", None)
    return combined


def _can_merge_structure_neighbors(left: dict, right: dict) -> bool:
    left_path = _section_path(left)
    right_path = _section_path(right)
    if left_path == right_path:
        return True
    if _is_formal_parent_boundary(left_path) or _is_formal_parent_boundary(right_path):
        return False
    if _section_structure_key(left) != _section_structure_key(right):
        return False
    return _compatible_numbered_paths(left_path, right_path)


def _is_formal_parent_boundary(path: str) -> bool:
    last = str(path or "").split(" / ")[-1].strip()
    return bool(_FORMAL_NUMBERED_HEADING_PATTERN.match(last))


def _compatible_numbered_paths(left_path: str, right_path: str) -> bool:
    left_num = _heading_number_tuple(left_path)
    right_num = _heading_number_tuple(right_path)
    if not left_num and not right_num:
        return True
    if not left_num or not right_num:
        return False
    if len(left_num) != len(right_num):
        return False
    return right_num >= left_num


def _heading_number_tuple(path: str) -> tuple[int, ...]:
    last = str(path or "").split(" / ")[-1].strip()
    match = re.match(r"^(\d+(?:\.\d+)*)\b", last)
    if not match:
        return ()
    return tuple(int(part) for part in match.group(1).split("."))


def _combine_section_label(left: str, right: str) -> str:
    if not right or right == left:
        return left
    if not left or left == "untitled":
        return right
    return f"{left} + {right}"


def _section_structure_key(section: dict) -> str:
    return _sect_key(_section_path(section))


def _section_path(section: dict) -> str:
    return str(section.get("section_path") or section.get("title") or "").strip()


def _sect_key(path: str) -> str:
    parts = [part.strip() for part in str(path or "").split(" / ") if part.strip()]
    if not parts:
        return "top:untitled"
    last = parts[-1]
    number = re.match(r"^(\d+)", last)
    if number:
        return f"num:{number.group(1)}"
    if len(parts) > 1:
        return f"h2:{parts[1].casefold()}"
    return f"top:{last.casefold()}"


def _is_atomic_section(section: dict) -> bool:
    return str(section.get("block_type") or "").strip().lower() in _ATOMIC_BLOCK_TYPES


def _raise_if_atomic_block_too_large(section: dict, paper_id: str) -> None:
    text = str(section.get("text") or "")
    max_size = settings.parent_chunk_max_size
    if _count_tokens(text) <= max_size:
        return
    if str(section.get("block_type") or "").strip().lower() == "table":
        return
    detail = _oversize_detail(_section_path(section), _atomic_block_label(section), text, max_size)
    _raise_chunking_error(_ERROR_STRUCTURED_BLOCK_TOO_LARGE, paper_id, detail)


def _atomic_block_label(section: dict) -> str:
    return str(section.get("block_type") or "structured block").strip()


def _raise_chunking_error(code: str, paper_id: str, detail: str) -> None:
    raise StructuredChunkingError(code, paper_id=paper_id or None, file_id=paper_id or None, detail=detail)


def _oversize_detail(section_path: str, unit_type: str, text: str, max_tokens: int) -> str:
    token_count = _count_tokens(text)
    location = section_path or "unknown section"
    return f"{unit_type} too large in {location}: tokens={token_count} limit={max_tokens}"


def _split_large_section(
    text: str, *, max_size: int, overlap_tokens: int = 0, paper_id: str = "", section_path: str = "",
) -> List[str]:
    """参数为章节正文、大小、重叠和定位；返回按自然边界切分的父块正文。"""
    return [part.text for part in _split_mapped_section(_MappedText(text), max_size=max_size,
        overlap_tokens=overlap_tokens, paper_id=paper_id, section_path=section_path)]


def _split_mapped_section(
    value: _MappedText, *, max_size: int, overlap_tokens: int, paper_id: str, section_path: str,
) -> list[_MappedText]:
    """参数为章节正文及预算/定位；返回携带精确来源的自然段父块，重叠同步裁剪来源。"""
    paragraphs = _split_mapped(value, r"\n{2,}")
    content_budget = _token_budget(max_size, overlap_tokens)
    parts: list[_MappedText] = []
    buffer = _MappedText("")
    for paragraph in paragraphs:
        candidate = _join_mapped([buffer, paragraph], _PARAGRAPH_SEPARATOR)
        if _count_tokens(candidate.text) <= content_budget:
            buffer = candidate
            continue
        if buffer.text:
            parts.append(buffer)
            buffer = _MappedText("")
        if _count_tokens(paragraph.text) <= content_budget:
            buffer = paragraph
            continue
        parts.extend(_split_mapped_paragraph(paragraph, content_budget, paper_id=paper_id, section_path=section_path))
    if buffer.text:
        parts.append(buffer)
    if _contains_formula(value.text):
        return parts
    return _overlap_mapped(parts, overlap_tokens, max_size, _PARAGRAPH_SEPARATOR)


def _split_oversize_paragraph(
    text: str, max_size: int, overlap_tokens: int = 0, *, paper_id: str = "", section_path: str = "",
) -> List[str]:
    """参数为超长段落及预算/定位；返回句子或完整公式边界切分结果，无法自然切分时显式失败。"""
    parts = _split_mapped_paragraph(_MappedText(text), _token_budget(max_size, overlap_tokens),
        paper_id=paper_id, section_path=section_path)
    if not _contains_formula(text):
        parts = _overlap_mapped(parts, overlap_tokens, max_size, _TOKEN_SPLITTER_SEPARATOR)
    return [part.text for part in parts]


def _split_mapped_paragraph(
    value: _MappedText, max_size: int, *, paper_id: str, section_path: str,
) -> list[_MappedText]:
    """参数为带来源段落和额度/定位；返回保持完整句子或公式的分块，拒绝单句强制裁剪。"""
    pieces: list[_MappedText] = []
    buffer = _MappedText("")
    for unit in _formula_safe_values(value):
        candidate = _join_mapped([buffer, unit], _TOKEN_SPLITTER_SEPARATOR)
        if _count_tokens(candidate.text) <= max_size:
            buffer = candidate
            continue
        if buffer.text:
            pieces.append(buffer)
            buffer = _MappedText("")
        if _count_tokens(unit.text) > max_size:
            label = "formula-safe unit" if _contains_formula(value.text) else "sentence"
            _raise_chunking_error(_ERROR_SENTENCE_TOO_LARGE, paper_id,
                _oversize_detail(section_path, label, unit.text, max_size))
        buffer = unit
    if buffer.text:
        pieces.append(buffer)
    return pieces


def _formula_safe_values(value: _MappedText) -> list[_MappedText]:
    """参数为带来源段落；返回句子和整段显示公式，所有切分位置来自原文正则坐标。"""
    units = []
    cursor = 0
    for match in _FORMULA_SPAN_PATTERN.finditer(value.text):
        prefix = value.slice(cursor, match.start()).strip()
        units.extend(_split_mapped(prefix, r"(?<=[.!?])\s+"))
        units.append(value.slice(match.start(), match.end()).strip())
        cursor = match.end()
    units.extend(_split_mapped(value.slice(cursor, len(value.text)), r"(?<=[.!?])\s+"))
    return units


def _contains_formula(text: str) -> bool:
    """参数为正文；返回是否存在完整公式，公式上下文不采用可能切半表达式的重叠。"""
    return _FORMULA_SPAN_PATTERN.search(str(text or "")) is not None


def _strip_markdown_for_embedding(text: str) -> str:
    """去除 Markdown 标记、图片占位等，产出纯文本用于向量化；公式片段整体保留。"""
    cleaned = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    cleaned = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", cleaned)
    cleaned = "".join(_strip_markdown_segment(part) for part in _FORMULA_SPAN_PATTERN.split(cleaned))
    return re.sub(r"\s+", " ", cleaned).strip()


def _strip_markdown_segment(segment: str, markers=_MARKDOWN_MARKERS_PATTERN) -> str:
    """参数为公式外的文本片段与标记集合；返回剥掉 Markdown 标记的文本，公式片段原样返回。"""
    if _FORMULA_SPAN_PATTERN.fullmatch(segment):
        return segment
    return markers.sub("", segment)


def _strip_formula_aware(text: str, *, keep_table_pipes: bool = False) -> str:
    """按公式感知方式剥 Markdown；表格索引文本额外保留列分隔符，同时保留换行。"""
    markers = _TABLE_MARKDOWN_MARKERS_PATTERN if keep_table_pipes else _MARKDOWN_MARKERS_PATTERN
    stripped = "".join(_strip_markdown_segment(part, markers) for part in _FORMULA_SPAN_PATTERN.split(str(text or "")))
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in stripped.splitlines()]
    return "\n".join(line for line in lines if line)


def _child_structure_units(value: _MappedText) -> list[dict[str, Any]]:
    """参数为父块正文；返回原子块、段落和枚举单元，保留原文的直接切片坐标。"""
    units = []
    for paragraph in _split_mapped(value, r"\n{2,}"):
        if _is_atomic_child_text(paragraph.text):
            units.append({"text": paragraph.text, "value": paragraph, "kind": "atomic"})
            continue
        matches = list(_unprotected_matches(paragraph.text, _ENUMERATED_UNIT_PATTERN))
        boundaries = [0, *(match.start() for match in matches if match.start() > 0), len(paragraph.text)]
        starts = {match.start() for match in matches}
        for start, end in zip(boundaries, boundaries[1:]):
            part = paragraph.slice(start, end).strip()
            if not part.text:
                continue
            kind = "enum" if start in starts else "paragraph"
            if _looks_like_markdown_list_item(part.text):
                kind = "list"
            units.append({"text": part.text, "value": part, "kind": kind})
    return units


def _is_atomic_child_text(text: str) -> bool:
    """参数为自然段；返回是否为必须保持完整的代码、公式或表格。"""
    stripped = str(text or "").strip()
    if stripped.startswith(("```", "~~~", "$$")):
        return True
    return _looks_like_markdown_table_line(stripped.splitlines()[0]) if stripped else False


def _pack_child_units(
    units: list[dict[str, Any]], *, target_size: int, paper_id: str, section_path: str,
) -> list[_MappedText]:
    """参数为结构单元、目标大小与定位；返回带来源的子块，结构兄弟单元保持既有边界。"""
    max_size = int(target_size * _CHILD_MAX_MULTIPLIER)
    min_size = max(1, min(_CHILD_SHORT_UNIT_TOKEN_LIMIT, int(target_size * _CHILD_SHORT_UNIT_RATIO)))
    groups: list[list[dict[str, Any]]] = []
    buffer: list[dict[str, Any]] = []
    for unit in units:
        structural_boundary = _is_structural_unit(unit) and any(_is_structural_unit(item) for item in buffer)
        candidate = _join_mapped((item["value"] for item in [*buffer, unit]), _TOKEN_SPLITTER_SEPARATOR)
        if buffer and ((structural_boundary and _count_tokens(unit["text"]) >= min_size)
                or _count_tokens(candidate.text) > max_size):
            groups.append(buffer)
            buffer = []
        buffer = [*buffer, unit]
    if buffer:
        groups.append(buffer)
    chunks = []
    for group in groups:
        value = _join_mapped((unit["value"] for unit in group), _TOKEN_SPLITTER_SEPARATOR)
        if _count_tokens(value.text) <= max_size:
            chunks.append(value)
        else:
            chunks.extend(_split_mapped_paragraph(value, target_size, paper_id=paper_id, section_path=section_path))
    return chunks


def _is_structural_unit(unit: dict[str, Any]) -> bool:
    """参数为子块结构单元；返回是否存在需保持的枚举、列表或原子结构边界。"""
    return unit.get("kind") in {"enum", "list", "atomic"} or bool(_STRUCTURAL_UNIT_PATTERN.match(unit.get("text", "")))


def _split_to_child_chunks(
    parent_text: str,
    parent_id: str,
    page_num: int,
    section: str,
    target_size: int,
    *,
    overlap_tokens: int = 0,
    block_type: str = "paragraph",
    section_path: str = "",
    asset_ref: str = "",
    parser_mode: str = "",
    paper_id: str = "",
    source_embedding_text: str = "",
    parent_source_fragments: tuple[SourceFragment, ...] = (),
) -> List[Chunk]:
    """参数为父正文、身份、结构和来源；返回同步携带子/父实际范围的分块，不复制父范围给子块。"""
    parent = _MappedText(parent_text, tuple(parent_source_fragments))
    units = _child_structure_units(parent)
    if not units:
        return []
    if block_type in _ATOMIC_BLOCK_TYPES:
        values = [parent]
    else:
        values = _pack_child_units(units, target_size=target_size, paper_id=paper_id, section_path=section_path)
        if not _contains_formula(parent_text) and not any(_is_structural_unit(unit) for unit in units):
            values = _overlap_mapped(values, overlap_tokens, int(target_size * _CHILD_MAX_MULTIPLIER),
                _TOKEN_SPLITTER_SEPARATOR)
    return [_child_chunk(
        parent_text=parent_text, child_text=value.text, parent_id=parent_id, child_idx=index,
        page_num=page_num, section=section, block_type=block_type, section_path=section_path,
        asset_ref=asset_ref, parser_mode=parser_mode, source_embedding_text=source_embedding_text,
        source_fragments=value.fragments, parent_source_fragments=parent.fragments)
        for index, value in enumerate(values)]


def _child_chunk(
    *,
    parent_text: str,
    child_text: str,
    parent_id: str,
    child_idx: int,
    page_num: int,
    section: str,
    block_type: str,
    section_path: str,
    asset_ref: str,
    parser_mode: str,
    source_embedding_text: str = "",
    source_fragments: tuple[SourceFragment, ...] = (),
    parent_source_fragments: tuple[SourceFragment, ...] = (),
) -> Chunk:
    """参数为子/父正文及结构来源；返回索引/阅读共用分块，多页不冒充首块单页来源。"""
    if source_fragments:
        pages = {fragment.get("page_num") for fragment in source_fragments}
        page_num = (next(iter(pages)) or 0) if len(pages) == 1 else 0
    chunk = Chunk(
        chunk_id=f"{parent_id}_{child_idx}",
        parent_id=parent_id,
        text=child_text,
        full_text=parent_text,
        page_num=page_num,
        section=section,
        embedding_text=_embedding_text(
            child_text, section_path,
            source_text=source_embedding_text,
            keep_table_pipes=_is_table_block(block_type),
        ),
        display_text=child_text,
        block_type=block_type,
        section_path=section_path,
        asset_ref=asset_ref,
        parser_mode=parser_mode,
        chunk_token_count=_count_tokens(child_text),
        full_text_token_count=_count_tokens(parent_text),
        source_fragments=list(source_fragments),
        parent_source_fragments=list(parent_source_fragments),
    )
    warnings = _chunk_warnings(parent_text, block_type)
    if warnings:
        chunk["chunk_warnings"] = warnings
    return chunk


def _chunk_warnings(parent_text: str, block_type: str) -> list[str]:
    if not _is_table_block(block_type):
        return []
    if _count_tokens(parent_text) <= settings.parent_chunk_max_size:
        return []
    return [_WARNING_ATOMIC_BLOCK_EXCEEDS_PARENT_LIMIT]


def _is_table_block(block_type: str) -> bool:
    """参数为块类型；返回该块是否为表格，表格的索引文本按列边界另行处理。"""
    return str(block_type or "").strip().lower() == "table"


def _embedding_text(text: str, section_path: str, *, source_text: str = "", keep_table_pipes: bool = False) -> str:
    """参数为子块文本、章节路径、可选来源索引文本与表格列分隔符开关；返回带头部的索引文本。"""
    # 表格的列分隔与行边界由块类型决定，没有替代来源文本时同样不能走会压平表格的通用清洗
    if keep_table_pipes or source_text:
        body = _strip_formula_aware(source_text or text, keep_table_pipes=keep_table_pipes)
    else:
        body = _strip_markdown_for_embedding(text)
    header = _strip_markdown_for_embedding(section_path)
    return f"{header}\n{body}".strip() if header else body


def _collect_pdf_sections(doc) -> List[dict]:
    """参数为受管 PDF 文档；返回含真实物理页码的文本片段，保留空白页的编号。"""
    sections: List[dict] = []
    for index in range(doc.page_count):
        text = doc.page_text(index).strip()
        if not text:
            continue
        title, body = _split_pdf_title_and_body(text)
        sections.append(_section(title, body, index + 1))
    return sections


def _split_pdf_title_and_body(text: str) -> tuple[str, str]:
    """尝试从 PDF 页首推断标题行和正文。"""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return "unknown", ""
    first_line = lines[0]
    # 页首第一行较短时，优先视作标题。
    if len(first_line) <= 120 and len(lines) > 1:
        return first_line, "\n".join(lines[1:]).strip()
    return "unknown", "\n".join(lines).strip()


def _read_text_file(file_path: str) -> str:
    """按预设编码顺序读取文本文件。"""
    data = Path(file_path).read_bytes()
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    # 所有编码都失败时，保底用 replace，尽量不让解析直接中断。
    return data.decode("utf-8", errors="replace")


def _strip_tex_comments(text: str) -> str:
    """删除 LaTeX 行内注释。"""
    return re.sub(r"(?<!\\)%.*$", "", text, flags=re.MULTILINE)


def _clean_latex(text: str) -> str:
    """去除常见 LaTeX 命令，保留较干净的正文文本。"""
    cleaned = re.sub(r"\\(?:begin|end)\{[^}]+\}", " ", text)
    cleaned = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?(?:\{[^}]*\})?", " ", cleaned)
    cleaned = re.sub(r"[{}]", " ", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def _table_to_markdown(rows: Iterable) -> str:
    """把表格行渲染成 Markdown 表格。"""
    rendered_rows = [_render_table_row(row) for row in rows]
    rendered_rows = [row for row in rendered_rows if row]
    if not rendered_rows:
        return ""
    if len(rendered_rows) == 1:
        return rendered_rows[0]
    # 多行表格时自动补 Markdown 分隔线。
    divider = "| " + " | ".join(["---"] * len(rendered_rows[0].split("|")[1:-1])) + " |"
    return "\n".join([rendered_rows[0], divider, *rendered_rows[1:]])


def _render_table_row(row) -> str:
    """渲染单行表格。"""
    cells = [cell.text.strip() for cell in row.cells]
    if not any(cells):
        return ""
    return "| " + " | ".join(cells) + " |"
